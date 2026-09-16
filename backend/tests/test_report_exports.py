"""V05-V06: polished server-side PDF and typed Excel exports."""

from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.shapes import Polygon

from app.reports.exports import _report_chart, render_excel_bytes, render_pdf_bytes
from app.reports.schema import ChartSpec


def _report_content() -> dict:
    return {
        "schemaVersion": "weekly-report/v1",
        "identity": {
            "reportId": "report-1",
            "versionId": "version-1",
            "title": "生产经营周报",
            "reportType": "weekly",
            "periodStart": "2026-08-03T00:00:00",
            "periodEnd": "2026-08-10T00:00:00",
            "comparisonPeriodStart": "2026-07-27T00:00:00",
            "comparisonPeriodEnd": "2026-08-03T00:00:00",
            "timezone": "Asia/Shanghai",
        },
        "managementSummary": {
            "findings": [
                {
                    "level": "confirmed",
                    "text": "本周总产量提升 6.12%，关键产线运行平稳。",
                    "evidenceArtifactIds": ["artifact-output"],
                    "evidenceStatus": "verified",
                    "confidence": 94,
                },
                {
                    "level": "attention",
                    "text": "AOI 工序不良率高于目标，需要在下周完成参数复核。",
                    "evidenceArtifactIds": ["artifact-quality"],
                    "evidenceStatus": "verified",
                    "confidence": 86,
                },
            ],
            "actions": [
                {
                    "text": "复核 AOI 工序参数并跟踪改善结果。",
                    "owner": "质量负责人",
                    "dueAt": "2026-08-14T18:00:00",
                    "priority": "high",
                    "evidenceArtifactIds": ["artifact-quality"],
                    "evidenceStatus": "verified",
                }
            ],
        },
        "metrics": [
            {
                "key": "output_qty",
                "name": "总产量",
                "value": 12840,
                "unit": "件",
                "previousValue": 12100,
                "changeRate": 0.0612,
                "trend": "up",
                "evidenceArtifactIds": ["artifact-output"],
                "confidence": 94,
            },
            {
                "key": "pass_rate",
                "name": "一次通过率",
                "value": 98.2,
                "unit": "%",
                "previousValue": 97.8,
                "changeRate": 0.0041,
                "trend": "up",
                "evidenceArtifactIds": ["artifact-quality"],
                "confidence": 91,
            },
            {
                "key": "downtime",
                "name": "非计划停机",
                "value": 4.6,
                "unit": "小时",
                "previousValue": 6.1,
                "changeRate": -0.2459,
                "trend": "down",
                "evidenceArtifactIds": ["artifact-equipment"],
                "confidence": 89,
            },
        ],
        "trends": [
            {
                "metricKey": "output_qty",
                "label": "近四周总产量",
                "unit": "件",
                "points": [
                    {"label": "W29", "value": 11520},
                    {"label": "W30", "value": 11880},
                    {"label": "W31", "value": 12100},
                    {"label": "W32", "value": 12840},
                ],
            }
        ],
        "anomalies": [
            {
                "title": "AOI 工序不良率偏高",
                "description": "周三夜班出现连续波动，建议复核贴片温度与检测阈值。",
                "severity": "high",
                "metricKey": "pass_rate",
                "evidenceArtifactIds": ["artifact-quality"],
                "evidenceStatus": "verified",
                "confidence": 86,
            }
        ],
        "progress": [
            {
                "title": "瓶颈工序改善",
                "status": "on_track",
                "note": "已完成两条产线节拍调整。",
                "completionRate": 72,
            }
        ],
        "recommendations": [
            {
                "title": "优先治理 AOI 波动",
                "rationale": "该工序对一次通过率贡献最大。",
                "priority": "high",
                "evidenceArtifactIds": ["artifact-quality"],
                "evidenceStatus": "verified",
            }
        ],
        "dataQuality": {
            "coverage": 96,
            "delayedSources": [],
            "missingMetrics": [],
            "temporaryMetrics": [],
            "failedSections": [],
        },
        "audience": {
            "scopeHash": "scope-a",
            "dataSourceIds": [1],
            "sensitiveFieldsRedacted": True,
        },
        "appendix": {
            "sources": ["MES 生产执行系统", "QMS 质量管理系统"],
            "definitions": [
                {"name": "一次通过率", "definition": "一次检验合格数量 / 总检验数量"}
            ],
        },
    }


def test_pdf_is_multi_page_searchable_and_contains_a_data_insight_image():
    payload = render_pdf_bytes(_report_content())

    assert payload.startswith(b"%PDF-")
    reader = PdfReader(BytesIO(payload))
    assert len(reader.pages) >= 2
    extracted = "".join(page.extract_text() or "" for page in reader.pages)
    assert "生产经营周报" in extracted
    assert "AOI 工序不良率偏高" in extracted
    resources = reader.pages[0]["/Resources"]
    xobjects = resources.get("/XObject", {})
    assert any(item.get_object().get("/Subtype") == "/Image" for item in xobjects.values())


def test_pdf_uses_the_same_agent_chart_spec_title_and_description_as_the_web():
    content = _report_content()
    content["charts"] = [
        {
            "key": "quality",
            "type": "pareto",
            "title": "缺陷类型 Pareto",
            "description": "按缺陷类型定位不良数量的主要贡献项。",
            "unit": "件",
            "categoryLabel": "缺陷类型",
            "valueLabel": "不良数量",
            "points": [
                {"label": "划伤", "value": 60, "secondaryValue": 60},
                {"label": "脏污", "value": 40, "secondaryValue": 100},
            ],
            "evidenceArtifactIds": ["artifact-quality"],
            "colorRole": "quality",
        }
    ]

    reader = PdfReader(BytesIO(render_pdf_bytes(content)))
    extracted = "".join(page.extract_text() or "" for page in reader.pages)

    assert "缺陷类型 Pareto" in extracted
    assert "按缺陷类型定位不良数量的主要贡献项" in extracted


def test_pdf_pareto_contains_bars_and_a_cumulative_curve():
    chart = _report_chart(ChartSpec.model_validate({
        "key": "quality", "type": "pareto", "title": "缺陷 Pareto",
        "description": "主要贡献项", "unit": "件", "categoryLabel": "缺陷",
        "valueLabel": "数量", "points": [
            {"label": "划伤", "value": 60, "secondaryValue": 60},
            {"label": "脏污", "value": 40, "secondaryValue": 100},
        ], "evidenceArtifactIds": ["artifact-quality"], "colorRole": "quality",
    }))

    assert any(isinstance(item, VerticalBarChart) for item in chart.contents)
    assert any(isinstance(item, HorizontalLineChart) for item in chart.contents)


def test_pdf_area_chart_has_a_visible_filled_region():
    chart = _report_chart(ChartSpec.model_validate({
        "key": "output", "type": "area", "title": "日产量趋势",
        "description": "产量变化", "unit": "件", "categoryLabel": "日期",
        "valueLabel": "产量", "points": [
            {"label": "08-08", "value": 100}, {"label": "08-09", "value": 120},
        ], "evidenceArtifactIds": ["artifact-output"], "colorRole": "production",
    }))

    assert any(isinstance(item, Polygon) for item in chart.contents)


def test_pdf_heatmap_keeps_missing_cells_blank_instead_of_zero():
    table = _report_chart(ChartSpec.model_validate({
        "key": "quality", "type": "heatmap", "title": "缺陷矩阵",
        "description": "交叉分布", "unit": "件", "categoryLabel": "工序",
        "valueLabel": "数量", "points": [
            {"label": "AOI", "group": "划伤", "value": 8},
            {"label": "SPI", "group": "脏污", "value": 5},
        ], "evidenceArtifactIds": ["artifact-quality"], "colorRole": "quality",
    }))

    flattened = [str(cell) for row in table._cellvalues for cell in row]
    assert flattened.count("—") == 2


def test_pdf_calendar_renders_a_week_grid_table():
    table = _report_chart(ChartSpec.model_validate({
        "key": "daily", "type": "calendar", "title": "日产量日历",
        "description": "逐日分布", "unit": "件", "categoryLabel": "日期",
        "valueLabel": "产量", "points": [
            {"label": "2026-06-01", "value": 100}, {"label": "2026-06-02", "value": 120},
            {"label": "2026-06-03", "value": 90}, {"label": "2026-06-04", "value": 140},
        ], "evidenceArtifactIds": ["artifact-output"], "colorRole": "production",
    }))
    assert table is not None
    assert len(table._cellvalues[0]) == 7  # 周一到周日表头
    # 四天都在同一周（2026-06-01 是周一），首周行应含四个数字单元格
    first_week = table._cellvalues[1]
    day_texts = [getattr(cell, "text", "") for cell in first_week]
    assert any(">1<" in text for text in day_texts)
    assert any(">4<" in text for text in day_texts)


def test_pdf_trend_compare_draws_two_line_series():
    chart = _report_chart(ChartSpec.model_validate({
        "key": "trend", "type": "line", "title": "产量趋势",
        "description": "两期对比", "unit": "件", "categoryLabel": "周",
        "valueLabel": "产量", "comparisonLabel": "上一周期", "points": [
            {"label": "2026-23", "value": 120, "previousValue": 90},
            {"label": "2026-24", "value": 135, "previousValue": 100},
        ], "evidenceArtifactIds": ["artifact-output"], "colorRole": "production",
    }))
    line = next(item for item in chart.contents if isinstance(item, HorizontalLineChart))
    assert len(line.data) == 2


def test_pdf_comparison_missing_values_are_gaps_for_lines_and_columns():
    content = _report_content()
    for kind, chart_class in (("line", HorizontalLineChart), ("column", VerticalBarChart)):
        spec = {"key": kind, "type": kind, "title": "双期产量", "description": "缺失不补零", "unit": "件", "categoryLabel": "日期", "valueLabel": "产量", "comparisonLabel": "上一周期", "points": [{"label": "2026-08-08", "value": 10, "previousValue": 8}, {"label": "2026-08-09", "value": 11, "previousValue": None}], "evidenceArtifactIds": ["artifact-output"], "colorRole": "production"}
        chart = next(item for item in _report_chart(ChartSpec.model_validate(spec)).contents if isinstance(item, chart_class))
        assert list(chart.data[1]) == [8, None]
        content["charts"] = [spec]
        assert render_pdf_bytes(content).startswith(b"%PDF")
        assert render_excel_bytes(content).startswith(b"PK")


def test_pdf_exports_every_agent_chart_instead_of_only_the_first_one():
    content = _report_content()
    content["charts"] = [
        {
            "key": "trend", "type": "line", "title": "日产量趋势",
            "description": "查看最近两日产量变化。", "unit": "件",
            "categoryLabel": "日期", "valueLabel": "产量",
            "points": [{"label": "08-08", "value": 100}, {"label": "08-09", "value": 120}],
            "evidenceArtifactIds": ["artifact-output"], "colorRole": "production",
        },
        {
            "key": "share", "type": "donut", "title": "产品产量构成",
            "description": "查看产品产量份额。", "unit": "件",
            "categoryLabel": "产品", "valueLabel": "产量",
            "points": [{"label": "A", "value": 70}, {"label": "B", "value": 30}],
            "evidenceArtifactIds": ["artifact-output"], "colorRole": "production",
        },
        {
            "key": "relationship", "type": "scatter", "title": "运行与停机关系",
            "description": "观察运行时长与停机时长的关系。", "unit": "小时",
            "secondaryUnit": "分钟", "categoryLabel": "设备", "valueLabel": "运行时长",
            "secondaryValueLabel": "停机时长",
            "points": [
                {"label": "A01", "value": 120, "secondaryValue": 8},
                {"label": "A02", "value": 150, "secondaryValue": 4},
            ],
            "evidenceArtifactIds": ["artifact-equipment"], "colorRole": "equipment",
        },
    ]

    reader = PdfReader(BytesIO(render_pdf_bytes(content)))
    extracted = "".join(page.extract_text() or "" for page in reader.pages)

    assert "日产量趋势" in extracted
    assert "产品产量构成" in extracted
    assert "运行与停机关系" in extracted
    assert "观察运行时长与停机时长的关系" in extracted


def test_pdf_bar_and_exact_table_keep_all_twelve_agent_points():
    content = _report_content()
    content["charts"] = [{
        "key": "lines", "type": "bar", "title": "产线产量排名",
        "description": "展示十二条产线。", "unit": "件",
        "categoryLabel": "产线", "valueLabel": "产量",
        "points": [
            {"label": f"产线-{index:02d}", "value": 130 - index}
            for index in range(1, 13)
        ],
        "evidenceArtifactIds": ["artifact-output"], "colorRole": "production",
    }]

    reader = PdfReader(BytesIO(render_pdf_bytes(content)))
    extracted = "".join(page.extract_text() or "" for page in reader.pages)

    assert "产线-11" in extracted
    assert "产线-12" in extracted


def test_excel_uses_required_sheets_and_real_numeric_cells():
    payload = render_excel_bytes(_report_content())

    workbook = load_workbook(BytesIO(payload), data_only=True)
    assert workbook.sheetnames == ["指标概览", "图表数据", "异常明细", "行动建议", "数据来源与口径"]
    overview = workbook["指标概览"]
    assert overview.freeze_panes == "A6"
    assert overview["B7"].value == 12840
    assert overview["B7"].data_type == "n"
    assert overview["F7"].value == 94
    assert workbook["异常明细"]["A7"].value == "AOI 工序不良率偏高"


def test_excel_exports_all_chart_data_and_native_visuals():
    content = _report_content()
    content["charts"] = [
        {
            "key": "trend", "type": "line", "title": "日产量趋势",
            "description": "查看产量变化。", "unit": "件", "categoryLabel": "日期", "valueLabel": "产量",
            "points": [{"label": "08-08", "value": 100}, {"label": "08-09", "value": 120}],
            "evidenceArtifactIds": ["artifact-output"], "colorRole": "production",
        },
        {
            "key": "share", "type": "donut", "title": "产品产量构成",
            "description": "查看产量份额。", "unit": "件", "categoryLabel": "产品", "valueLabel": "产量",
            "points": [{"label": "A", "value": 70}, {"label": "B", "value": 30}],
            "evidenceArtifactIds": ["artifact-output"], "colorRole": "production",
        },
        {
            "key": "relationship", "type": "scatter", "title": "运行与停机关系",
            "description": "观察两个指标的关系。", "unit": "小时", "secondaryUnit": "分钟",
            "categoryLabel": "设备", "valueLabel": "运行时长", "secondaryValueLabel": "停机时长",
            "points": [
                {"label": "A01", "value": 120, "secondaryValue": 8},
                {"label": "A02", "value": 150, "secondaryValue": 4},
            ],
            "evidenceArtifactIds": ["artifact-equipment"], "colorRole": "equipment",
        },
    ]

    workbook = load_workbook(BytesIO(render_excel_bytes(content)), data_only=False)
    sheet = workbook["图表数据"]

    assert len(sheet._charts) == 3
    assert [sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)].count("日产量趋势") == 1
    assert [sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)].count("产品产量构成") == 1
    assert [sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)].count("运行与停机关系") == 1
    numeric_values = [sheet.cell(row, 3).value for row in range(1, sheet.max_row + 1)]
    assert 120 in numeric_values


def test_excel_dates_and_print_layout_are_readable_in_desktop_excel():
    workbook = load_workbook(BytesIO(render_excel_bytes(_report_content())))

    for sheet in workbook.worksheets:
        assert sheet["B3"].number_format == "yyyy-mm-dd"
        assert sheet["D3"].number_format == "yyyy-mm-dd"
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.sheet_properties.pageSetUpPr.fitToPage is True


def test_long_pdf_table_repeats_cleanly_and_keeps_last_row_searchable():
    content = _report_content()
    content["anomalies"] = [
        {
            "title": f"异常记录 {index:03d}",
            "description": "这是用于验证跨页表格、自动行高和长文本换行的异常说明。" * 2,
            "severity": "medium",
            "metricKey": "pass_rate",
            "evidenceArtifactIds": [f"artifact-{index}"],
            "evidenceStatus": "verified",
            "confidence": 80,
        }
        for index in range(1, 61)
    ]

    reader = PdfReader(BytesIO(render_pdf_bytes(content)))
    extracted = "".join(page.extract_text() or "" for page in reader.pages)

    assert len(reader.pages) >= 5
    assert "异常记录 060" in extracted


def test_excel_neutralizes_formula_injection_in_untrusted_text():
    content = _report_content()
    content["anomalies"][0]["title"] = '=HYPERLINK("https://evil.invalid","click")'

    workbook = load_workbook(BytesIO(render_excel_bytes(content)), data_only=False)
    cell = workbook["异常明细"]["A7"]

    assert cell.data_type == "s"
    assert cell.value.startswith("'=HYPERLINK")


def test_pdf_escapes_reportlab_markup_and_does_not_embed_local_files(tmp_path):
    local_image = tmp_path / "private.png"
    local_image.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDAT\x08\xd7c\xf8\xcf\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99=\x1d\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    content = _report_content()
    content["managementSummary"]["findings"][0]["text"] = (
        f'<img src="{local_image}"/> 本周总产量提升 6.12%。'
    )

    reader = PdfReader(BytesIO(render_pdf_bytes(content)))
    images = []
    for page in reader.pages:
        resources = page["/Resources"]
        images.extend(
            item.get_object()
            for item in resources.get("/XObject", {}).values()
            if item.get_object().get("/Subtype") == "/Image"
        )

    assert len(images) == 1  # only the renderer-owned insight hero
    assert "img" in "".join(page.extract_text() or "" for page in reader.pages)


def test_data_anomaly_pdf_uses_an_explicit_no_trend_state():
    content = _report_content()
    content["metrics"] = []
    content["trends"] = []
    content["dataQuality"]["coverage"] = 42

    reader = PdfReader(BytesIO(render_pdf_bytes(content)))
    extracted = "".join(page.extract_text() or "" for page in reader.pages)

    assert "当前冻结版本没有可验证的时间序列" in extracted
