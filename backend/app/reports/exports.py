"""Deterministic PDF and Excel renderers for frozen weekly-report content."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

import reportlab
from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PillowImage
from PIL import ImageDraw, ImageFont
from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.charts.lineplots import ScatterPlot
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing, Polygon, String
from reportlab.graphics.widgets.markers import makeMarker
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    Image,
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.reports.schema import ChartSpec, StructuredWeeklyReport, calculate_confidence

INK = HexColor("#172033")
MUTED = HexColor("#667085")
BRAND = HexColor("#0B6E69")
BRAND_DARK = HexColor("#084C49")
MINT = HexColor("#DDF4EF")
AMBER = HexColor("#F2A93B")
SURFACE = HexColor("#F5F7FA")
LINE = HexColor("#D9E0E8")


def _register_fonts() -> tuple[str, str]:
    regular = "ForgeSightSong"
    if regular not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light", "STSong-Light"


def _pdf_text(value: object) -> str:
    """Treat every model/user value as text, never as ReportLab markup."""

    return escape(str(value), {'"': "&quot;", "'": "&#39;"})


def _styles():
    regular, bold = _register_fonts()
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "ForgeSightTitle",
            parent=styles["Title"],
            fontName=bold,
            fontSize=24,
            leading=32,
            textColor=INK,
            alignment=TA_LEFT,
            spaceAfter=8,
        ),
        "eyebrow": ParagraphStyle(
            "ForgeSightEyebrow",
            parent=styles["Normal"],
            fontName=regular,
            fontSize=9,
            leading=12,
            textColor=BRAND,
            spaceAfter=4,
        ),
        "heading": ParagraphStyle(
            "ForgeSightHeading",
            parent=styles["Heading2"],
            fontName=bold,
            fontSize=16,
            leading=22,
            textColor=INK,
            spaceBefore=8,
            spaceAfter=8,
        ),
        "subheading": ParagraphStyle(
            "ForgeSightSubheading",
            parent=styles["Heading3"],
            fontName=bold,
            fontSize=11,
            leading=16,
            textColor=INK,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "ForgeSightBody",
            parent=styles["BodyText"],
            fontName=regular,
            fontSize=9.5,
            leading=15,
            textColor=INK,
        ),
        "small": ParagraphStyle(
            "ForgeSightSmall",
            parent=styles["BodyText"],
            fontName=regular,
            fontSize=7.5,
            leading=11,
            textColor=MUTED,
        ),
        "tableHeader": ParagraphStyle(
            "ForgeSightTableHeader",
            parent=styles["BodyText"],
            fontName=regular,
            fontSize=7.5,
            leading=11,
            textColor=colors.white,
        ),
        "metric": ParagraphStyle(
            "ForgeSightMetric",
            parent=styles["Normal"],
            fontName=bold,
            fontSize=17,
            leading=21,
            textColor=BRAND_DARK,
            alignment=TA_CENTER,
        ),
        "metricLabel": ParagraphStyle(
            "ForgeSightMetricLabel",
            parent=styles["Normal"],
            fontName=regular,
            fontSize=8,
            leading=11,
            textColor=MUTED,
            alignment=TA_CENTER,
        ),
    }


def _insight_image(report: StructuredWeeklyReport) -> BytesIO:
    """Create a high-resolution, deterministic insight hero from report data."""

    width, height = 1440, 480
    image = PillowImage.new("RGB", (width, height), "#0A4F4B")
    draw = ImageDraw.Draw(image)
    font_root = Path(reportlab.__file__).resolve().parent / "fonts"
    label_font = ImageFont.truetype(str(font_root / "Vera.ttf"), 20)
    headline_font = ImageFont.truetype(str(font_root / "VeraBd.ttf"), 58)
    caption_font = ImageFont.truetype(str(font_root / "Vera.ttf"), 16)
    for y in range(height):
        ratio = y / height
        color = (
            int(10 + 6 * ratio),
            int(79 + 35 * ratio),
            int(75 + 34 * ratio),
        )
        draw.line((0, y, width, y), fill=color)
    draw.rounded_rectangle((48, 42, width - 48, height - 42), radius=30, outline="#5AD6C8", width=3)
    confidence = calculate_confidence(report.data_quality.model_dump(by_alias=True))
    anomaly_variant = confidence["templateVariant"] == "data_anomaly"
    hero_label = "FORGESIGHT  /  DATA QUALITY ALERT" if anomaly_variant else "FORGESIGHT  /  WEEKLY SIGNAL"
    draw.text((82, 74), hero_label, fill="#B8F2E9", font=label_font)
    primary = report.metrics[0] if report.metrics and not anomaly_variant else None
    headline = f"{primary.value:,.0f}" if primary else f"{confidence['score']}%"
    subtitle = "PRIMARY KPI" if primary else "REPORT CONFIDENCE"
    draw.text((84, 145), headline, fill="white", font=headline_font)
    draw.text((86, 220), subtitle, fill="#B8F2E9", font=caption_font)

    chart_left, chart_top, chart_right, chart_bottom = 610, 92, 1340, 370
    draw.line((chart_left, chart_bottom, chart_right, chart_bottom), fill="#77CFC4", width=2)
    trend = report.trends[0] if report.trends and len(report.trends[0].points) >= 2 else None
    if trend:
        points = [point.value for point in trend.points]
        low, high = min(points), max(points)
        spread = high - low or 1
        coordinates = []
        for index, value in enumerate(points):
            x = chart_left + (chart_right - chart_left) * index / (len(points) - 1)
            y = chart_bottom - (chart_bottom - chart_top) * (value - low) / spread
            coordinates.append((x, y))
        draw.line(coordinates, fill="#F8C76A", width=8, joint="curve")
        for x, y in coordinates:
            draw.ellipse((x - 9, y - 9, x + 9, y + 9), fill="#F8C76A", outline="white", width=3)
    draw.text(
        (chart_left, 395),
        "TREND / VERIFIED DATA" if trend else "NO VERIFIED TREND DATA",
        fill="#B8F2E9",
        font=caption_font,
    )
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    output.seek(0)
    return output


def _page_decoration(canvas, document) -> None:
    regular, _ = _register_fonts()
    canvas.saveState()
    width, height = A4
    canvas.setFillColor(BRAND)
    canvas.rect(0, height - 7 * mm, width, 7 * mm, fill=1, stroke=0)
    canvas.setFont(regular, 7.5)
    canvas.setFillColor(MUTED)
    canvas.drawString(18 * mm, 10 * mm, "ForgeSight · 企业数据洞察")
    canvas.drawRightString(width - 18 * mm, 10 * mm, f"{canvas.getPageNumber():02d}")
    canvas.restoreState()


def _metric_table(report: StructuredWeeklyReport, styles: dict) -> Table:
    cells = []
    for metric in report.metrics[:4]:
        delta = ""
        if metric.change_rate is not None:
            delta = f"{metric.change_rate:+.1%} 环比"
        cells.append(
            [
                Paragraph(_pdf_text(metric.name), styles["metricLabel"]),
                Paragraph(_pdf_text(f"{metric.value:,.2f} {metric.unit}"), styles["metric"]),
                Paragraph(delta or f"可信度 {metric.confidence}", styles["metricLabel"]),
            ]
        )
    if not cells:
        cells.append(
            [
                Paragraph("数据覆盖率", styles["metricLabel"]),
                Paragraph(f"{report.data_quality.coverage}%", styles["metric"]),
                Paragraph("本期数据质量", styles["metricLabel"]),
            ]
        )
    table = Table([cells], colWidths=[170 * mm / len(cells)] * len(cells), rowHeights=[28 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), SURFACE),
                ("BOX", (0, 0), (-1, -1), 0.6, LINE),
                ("INNERGRID", (0, 0), (-1, -1), 0.6, LINE),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    return table


def _export_chart_specs(report: StructuredWeeklyReport) -> list[ChartSpec]:
    if report.charts:
        return list(report.charts)
    return [
        ChartSpec(
            key=f"legacy-{index + 1}",
            type="line",
            title=series.label,
            description=f"展示{series.label}的周期变化。",
            unit=series.unit,
            categoryLabel="周期",
            valueLabel=series.label,
            points=[point.model_dump(by_alias=True) for point in series.points],
            evidenceArtifactIds=[],
            colorRole="neutral",
        )
        for index, series in enumerate(report.trends)
        if len(series.points) >= 2
    ]


def _pdf_chart_accent(spec: ChartSpec):
    accents = {
        "ocean": HexColor("#25677A"),
        "teal": BRAND,
        "violet": HexColor("#7357C8"),
        "graphite": HexColor("#58636D"),
    }
    if spec.accent in accents:
        return accents[spec.accent]
    return AMBER if spec.color_role == "quality" else BRAND


def _parse_calendar_date(label: str):
    text = str(label).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            from datetime import datetime as _dt
            return _dt.strptime(text[:19], fmt).date()
        except ValueError:
            continue
    return None


def _calendar_heat_color(base, ratio: float):
    ratio = max(0.0, min(1.0, ratio))
    return colors.Color(
        base.red + (1 - base.red) * (1 - ratio),
        base.green + (1 - base.green) * (1 - ratio),
        base.blue + (1 - base.blue) * (1 - ratio),
    )


def _pdf_calendar(spec: ChartSpec) -> Table | None:
    values = []
    for point in spec.points:
        day = _parse_calendar_date(point.label)
        if day is None:
            continue
        values.append((day, point.value or 0))
    if len(values) < 2:
        return None
    values.sort(key=lambda item: item[0])
    low = min(value for _, value in values)
    high = max(value for _, value in values)
    spread = (high - low) or 1
    by_day = dict(values)
    base = _pdf_chart_accent(spec)
    weekday_labels = ["一", "二", "三", "四", "五", "六", "日"]
    first = values[0][0]
    last = values[-1][0]
    # align grid to Monday of the first week
    from datetime import timedelta
    grid_start = first - timedelta(days=first.weekday())
    total_days = (last - grid_start).days + 1
    weeks = (total_days + 6) // 7
    header = [Paragraph(f"<font color='#8997a3'>{label}</font>", _styles()["small"]) for label in weekday_labels]
    rows = [header]
    styles_cmds = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]
    for week in range(weeks):
        week_row = []
        for weekday in range(7):
            current = grid_start + timedelta(days=week * 7 + weekday)
            if current in by_day:
                ratio = (by_day[current] - low) / spread
                cell_color = _calendar_heat_color(base, ratio)
                text_hex = "#FFFFFF" if ratio > 0.55 else "#182329"
                week_row.append(Paragraph(f"<font color='{text_hex}'>{current.day}</font>", _styles()["small"]))
                styles_cmds.append(("BACKGROUND", (weekday, week + 1), (weekday, week + 1), cell_color))
            else:
                week_row.append("")
        while len(week_row) < 7:
            week_row.append("")
        rows.append(week_row)
    widths = [16 * mm] * 7
    table = Table(rows, colWidths=widths, rowHeights=[6 * mm] + [11 * mm] * weeks)
    table.setStyle(TableStyle(styles_cmds))
    return table


def _report_chart(spec: ChartSpec) -> Drawing | Table | None:
    points = spec.points
    if len(points) < 2:
        return None
    if spec.type == "table":
        return None
    if spec.type == "calendar":
        return _pdf_calendar(spec)
    drawing = Drawing(480, 190)
    if spec.type == "donut":
        chart = Pie()
        chart.x = 115
        chart.y = 8
        chart.width = 250
        chart.height = 165
        chart.data = [max(point.value, 0) for point in points[:8]]
        chart.labels = [point.label for point in points[:8]]
        chart.innerRadiusFraction = 0.55
        chart.slices.strokeColor = colors.white
        chart.slices.strokeWidth = 1
        palette = [_pdf_chart_accent(spec), AMBER, HexColor("#7357C8"), HexColor("#2B8A9A"), HexColor("#D95D5D")]
        for index in range(len(chart.data)):
            chart.slices[index].fillColor = palette[index % len(palette)]
            chart.slices[index].fontName = _register_fonts()[0]
            chart.slices[index].fontSize = 7
        drawing.add(chart)
        return drawing
    if spec.type == "scatter":
        usable = [
            (point.value, point.secondary_value)
            for point in points
            if point.secondary_value is not None
        ]
        if len(usable) < 2:
            return None
        chart = ScatterPlot()
        chart.x = 58
        chart.y = 34
        chart.height = 120
        chart.width = 370
        chart.data = [usable]
        chart.xValueAxis.labels.fontName = "Helvetica"
        chart.xValueAxis.labels.fontSize = 7
        chart.yValueAxis.labels.fontName = "Helvetica"
        chart.yValueAxis.labels.fontSize = 7
        chart.xValueAxis.gridStrokeColor = LINE
        chart.yValueAxis.gridStrokeColor = LINE
        chart.lines[0].strokeColor = None
        chart.lines[0].symbol = makeMarker("FilledCircle")
        chart.lines[0].symbol.fillColor = HexColor("#7357C8")
        chart.lines[0].symbol.strokeColor = colors.white
        chart.lines[0].symbol.size = 7
        drawing.add(chart)
        return drawing
    if spec.type == "heatmap":
        groups = list(dict.fromkeys(point.group for point in points if point.group))
        labels = list(dict.fromkeys(point.label for point in points))
        if not groups:
            return None
        values = {(point.label, point.group): point.value for point in points}
        all_values = [point.value for point in points]
        low, high = min(all_values), max(all_values)
        spread = high - low or 1
        rows = [[""] + [Paragraph(_pdf_text(group), _styles()["small"]) for group in groups]]
        for label in labels:
            rows.append(
                [Paragraph(_pdf_text(label), _styles()["small"])]
                + [
                    "—" if values.get((label, group)) is None else f"{values[(label, group)]:,.2f}"
                    for group in groups
                ]
            )
        widths = [32 * mm] + [min(28 * mm, 138 * mm / len(groups))] * len(groups)
        table = Table(rows, colWidths=widths, rowHeights=8 * mm)
        styles = [
            ("BACKGROUND", (0, 0), (-1, 0), BRAND_DARK),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("BACKGROUND", (0, 1), (0, -1), SURFACE),
            ("GRID", (0, 0), (-1, -1), 0.5, LINE),
            ("ALIGN", (1, 1), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (1, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (1, 1), (-1, -1), 7),
        ]
        for row_index, label in enumerate(labels, start=1):
            for column_index, group in enumerate(groups, start=1):
                value = values.get((label, group))
                if value is None:
                    styles.append(("BACKGROUND", (column_index, row_index), (column_index, row_index), SURFACE))
                    continue
                ratio = (value - low) / spread
                color = colors.Color(0.87 - ratio * 0.55, 0.96 - ratio * 0.42, 0.94 - ratio * 0.48)
                styles.append(("BACKGROUND", (column_index, row_index), (column_index, row_index), color))
        table.setStyle(TableStyle(styles))
        return table
    if spec.type == "combo" and all(point.secondary_value is not None for point in points):
        categories = [point.label for point in points]
        bar = VerticalBarChart()
        bar.x = 38
        bar.y = 30
        bar.height = 115
        bar.width = 175
        bar.data = [[point.value for point in points]]
        bar.categoryAxis.categoryNames = categories
        bar.categoryAxis.labels.fontName = _register_fonts()[0]
        bar.categoryAxis.labels.fontSize = 6
        bar.valueAxis.labels.fontName = "Helvetica"
        bar.valueAxis.labels.fontSize = 7
        bar.valueAxis.gridStrokeColor = LINE
        bar.bars[0].fillColor = _pdf_chart_accent(spec)
        bar.bars[0].strokeColor = None
        line = HorizontalLineChart()
        line.x = 270
        line.y = 30
        line.height = 115
        line.width = 175
        line.data = [[point.secondary_value for point in points]]
        line.categoryAxis.categoryNames = categories
        line.categoryAxis.labels.fontName = _register_fonts()[0]
        line.categoryAxis.labels.fontSize = 6
        line.valueAxis.labels.fontName = "Helvetica"
        line.valueAxis.labels.fontSize = 7
        line.valueAxis.gridStrokeColor = LINE
        line.lines[0].strokeColor = AMBER
        line.lines[0].strokeWidth = 2.5
        line.lines[0].symbol = makeMarker("FilledCircle")
        drawing.add(bar)
        drawing.add(line)
        drawing.add(String(38, 162, f"{spec.value_label} ({spec.unit})", fontName=_register_fonts()[0], fontSize=8, fillColor=BRAND_DARK))
        drawing.add(
            String(
                270,
                162,
                f"{spec.secondary_value_label or '次指标'} ({spec.secondary_unit or spec.unit})",
                fontName=_register_fonts()[0],
                fontSize=8,
                fillColor=AMBER,
            )
        )
        return drawing
    if spec.type == "pareto" and all(point.secondary_value is not None for point in points):
        visible = points[:12]
        categories = [point.label for point in visible]
        bar = VerticalBarChart()
        bar.x = 38
        bar.y = 30
        bar.height = 115
        bar.width = 175
        bar.data = [[point.value for point in visible]]
        bar.categoryAxis.categoryNames = categories
        bar.categoryAxis.labels.fontName = _register_fonts()[0]
        bar.categoryAxis.labels.fontSize = 6
        bar.valueAxis.labels.fontName = "Helvetica"
        bar.valueAxis.labels.fontSize = 7
        bar.valueAxis.gridStrokeColor = LINE
        bar.bars[0].fillColor = _pdf_chart_accent(spec)
        bar.bars[0].strokeColor = None
        cumulative = HorizontalLineChart()
        cumulative.x = 270
        cumulative.y = 30
        cumulative.height = 115
        cumulative.width = 175
        cumulative.data = [[point.secondary_value for point in visible]]
        cumulative.categoryAxis.categoryNames = categories
        cumulative.categoryAxis.labels.fontName = _register_fonts()[0]
        cumulative.categoryAxis.labels.fontSize = 6
        cumulative.valueAxis.valueMin = 0
        cumulative.valueAxis.valueMax = 100
        cumulative.valueAxis.valueStep = 20
        cumulative.valueAxis.labels.fontName = "Helvetica"
        cumulative.valueAxis.labels.fontSize = 7
        cumulative.valueAxis.gridStrokeColor = LINE
        cumulative.lines[0].strokeColor = BRAND
        cumulative.lines[0].strokeWidth = 2.5
        cumulative.lines[0].symbol = makeMarker("FilledCircle")
        drawing.add(bar)
        drawing.add(cumulative)
        drawing.add(String(38, 162, f"{spec.value_label} ({spec.unit})", fontName=_register_fonts()[0], fontSize=8, fillColor=AMBER))
        drawing.add(String(270, 162, "累计贡献率 (%)", fontName=_register_fonts()[0], fontSize=8, fillColor=BRAND_DARK))
        return drawing
    if spec.type == "column":
        chart = VerticalBarChart()
        chart.x = 55
        chart.y = 28
        chart.height = 132
        chart.width = 390
        chart.data = [[point.value for point in points[:16]]]
        if any(point.previous_value is not None for point in points[:16]):
            chart.data.append([point.previous_value for point in points[:16]])
            chart.bars[1].fillColor = colors.HexColor("#9aa8ac")
            chart.bars[1].strokeColor = None
            drawing.add(String(58, 168, "本期 / 上一周期（缺失值留空）", fontName=_register_fonts()[0], fontSize=7, fillColor=BRAND_DARK))
        chart.categoryAxis.categoryNames = [point.label for point in points[:16]]
        chart.categoryAxis.labels.fontName = _register_fonts()[0]
        chart.categoryAxis.labels.fontSize = 7
        chart.valueAxis.labels.fontName = "Helvetica"
        chart.valueAxis.labels.fontSize = 7
        chart.bars[0].fillColor = _pdf_chart_accent(spec)
        chart.bars[0].strokeColor = None
        chart.valueAxis.gridStrokeColor = LINE
        chart.valueAxis.gridStrokeDashArray = (2, 2)
        drawing.add(chart)
        return drawing
    if spec.type == "bar":
        chart = HorizontalBarChart()
        chart.x = 80
        chart.y = 28
        chart.height = 132
        chart.width = 350
        chart.data = [[point.value for point in reversed(points[:12])]]
        chart.categoryAxis.categoryNames = [point.label for point in reversed(points[:12])]
        chart.categoryAxis.labels.fontName = _register_fonts()[0]
        chart.categoryAxis.labels.fontSize = 7
        chart.valueAxis.labels.fontName = "Helvetica"
        chart.valueAxis.labels.fontSize = 7
        chart.bars[0].fillColor = _pdf_chart_accent(spec)
        chart.bars[0].strokeColor = None
        chart.valueAxis.gridStrokeColor = LINE
        chart.valueAxis.gridStrokeDashArray = (2, 2)
        drawing.add(chart)
        return drawing
    if spec.type == "area":
        values = [point.value for point in points]
        maximum = max(max(values), 0) or 1
        minimum = min(min(values), 0)
        spread = maximum - minimum or 1
        left, bottom, width, height = 55, 34, 390, 120
        step = width / max(len(points) - 1, 1)
        baseline = bottom + (0 - minimum) / spread * height
        polygon_points = [left, baseline]
        for index, value in enumerate(values):
            polygon_points.extend([left + index * step, bottom + (value - minimum) / spread * height])
        polygon_points.extend([left + width, baseline])
        drawing.add(Polygon(polygon_points, fillColor=MINT, strokeColor=None))
    has_previous = any(point.previous_value is not None for point in points)
    chart = HorizontalLineChart()
    chart.x = 55
    chart.y = 34
    chart.height = 120
    chart.width = 390
    series: list[list[float]] = [[point.value for point in points]]
    if has_previous:
        series.append([point.previous_value for point in points])
    chart.data = series
    chart.categoryAxis.categoryNames = [point.label for point in points]
    chart.categoryAxis.labels.fontName = _register_fonts()[0]
    chart.categoryAxis.labels.fontSize = 7
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 7
    chart.lines[0].strokeColor = _pdf_chart_accent(spec)
    chart.lines[0].strokeWidth = 2.5
    chart.lines[0].symbol = None
    if has_previous:
        chart.lines[1].strokeColor = HexColor("#98a4b0")
        chart.lines[1].strokeWidth = 1.6
        chart.lines[1].strokeDashArray = (4, 3)
        chart.lines[1].symbol = None
    chart.valueAxis.gridStrokeColor = LINE
    chart.valueAxis.gridStrokeDashArray = (2, 2)
    if spec.type == "area" and min(point.value for point in points) >= 0:
        chart.valueAxis.valueMin = 0
    drawing.add(chart)
    if has_previous:
        drawing.add(String(58, 168, "— 本期    ⋯ 上一周期", fontName=_register_fonts()[0], fontSize=7, fillColor=colors.HexColor("#5d6d72")))
    return drawing


def _chart_data_table(spec: ChartSpec, styles: dict) -> Table:
    has_secondary = any(point.secondary_value is not None for point in spec.points)
    has_previous = any(point.previous_value is not None for point in spec.points)
    secondary_label = "累计贡献率" if spec.type == "pareto" else (spec.secondary_value_label or "次指标")
    headers = [spec.category_label, f"{spec.value_label}（{spec.unit}）"]
    if spec.type == "table" and spec.composition_allowed:
        headers += ["占比", "累计占比"]
    if has_previous:
        headers.append(f"{spec.comparison_label or '上一周期'}（{spec.unit}）")
    if has_secondary:
        headers.append(f"{secondary_label}（{'%' if spec.type == 'pareto' else spec.secondary_unit or spec.unit}）")
    rows = [[Paragraph(_pdf_text(item), styles["tableHeader"]) for item in headers]]
    table_total = sum(max(point.value or 0, 0) for point in spec.points) or 1
    table_cumulative = 0.0
    for point in spec.points[:12]:
        label = point.label if not point.group else f"{point.label} / {point.group}"
        row = [Paragraph(_pdf_text(label), styles["small"]), f"{point.value:,.2f}"]
        if spec.type == "table" and spec.composition_allowed:
            share = max(point.value or 0, 0) / table_total * 100
            table_cumulative += share
            row += [f"{share:.1f}%", f"{min(table_cumulative, 100.0):.1f}%"]
        if has_previous:
            row.append("—" if point.previous_value is None else f"{point.previous_value:,.2f}")
        if has_secondary:
            row.append("—" if point.secondary_value is None else f"{point.secondary_value:,.2f}")
        rows.append(row)
    widths = [72 * mm, 48 * mm] + ([26 * mm, 30 * mm] if spec.type == "table" and spec.composition_allowed else []) + ([42 * mm] if has_previous else []) + ([50 * mm] if has_secondary else [])
    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(_pdf_table_style())
    return table


def render_pdf_bytes(content: dict) -> bytes:
    """Render an A4, searchable PDF without a browser session."""

    report = StructuredWeeklyReport.model_validate(content)
    styles = _styles()
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=17 * mm,
        bottomMargin=18 * mm,
        title=report.identity.title,
        author="ForgeSight",
        subject="企业经营周报",
    )
    period = f"{report.identity.period_start:%Y.%m.%d} — {report.identity.period_end:%Y.%m.%d}"
    story = [
        Spacer(1, 7 * mm),
        Paragraph("FORGESIGHT · WEEKLY BUSINESS REVIEW", styles["eyebrow"]),
        Paragraph(_pdf_text(report.identity.title), styles["title"]),
        Paragraph(
            _pdf_text(
                f"报告周期 {period}　·　数据覆盖率 {report.data_quality.coverage}%　·　版本 {report.identity.version_id}"
            ),
            styles["small"],
        ),
        Spacer(1, 6 * mm),
    ]
    hero = Image(_insight_image(report), width=170 * mm, height=56.7 * mm)
    hero.altText = "根据本期真实指标生成的数据洞察主图"
    story.extend([hero, Spacer(1, 6 * mm), _metric_table(report, styles), Spacer(1, 6 * mm)])
    story.append(Paragraph("管理摘要", styles["heading"]))
    for finding in report.management_summary.findings[:5]:
        prefix = "●" if finding.level == "confirmed" else "◆"
        story.append(
            Paragraph(
                f"<font color='#0B6E69'>{prefix}</font>　{_pdf_text(finding.text)}　"
                f"<font color='#667085'>可信度 {finding.confidence}</font>",
                styles["body"],
            )
        )
        story.append(Spacer(1, 2.5 * mm))
    story.append(PageBreak())
    story.extend(
        [
            Paragraph("指标趋势与经营信号", styles["heading"]),
            Paragraph("图表仅使用已冻结版本中的结构化数值，确保与站内展示一致。", styles["small"]),
        ]
    )
    chart_specs = _export_chart_specs(report)
    if chart_specs:
        for index, chart_spec in enumerate(chart_specs):
            if index:
                story.append(PageBreak())
            story.extend(
                [
                    Spacer(1, 3 * mm),
                    Paragraph(_pdf_text(chart_spec.title), styles["subheading"]),
                    Paragraph(_pdf_text(chart_spec.description), styles["small"]),
                ]
            )
            chart = _report_chart(chart_spec)
            if chart is not None:
                story.extend([Spacer(1, 3 * mm), chart, Spacer(1, 2 * mm)])
            story.append(_chart_data_table(chart_spec, styles))
    else:
        story.extend(
            [
                Spacer(1, 3 * mm),
                Paragraph("当前冻结版本没有可验证的时间序列，因此未绘制趋势图。", styles["body"]),
                Spacer(1, 4 * mm),
            ]
        )

    story.append(PageBreak())
    story.append(Paragraph("异常明细", styles["heading"]))
    anomaly_rows = [
        [
            Paragraph("异常", styles["tableHeader"]),
            Paragraph("等级", styles["tableHeader"]),
            Paragraph("影响与说明", styles["tableHeader"]),
            Paragraph("可信度", styles["tableHeader"]),
        ]
    ]
    for item in report.anomalies:
        anomaly_rows.append(
            [
                Paragraph(_pdf_text(item.title), styles["body"]),
                Paragraph(item.severity.upper(), styles["small"]),
                Paragraph(_pdf_text(item.description), styles["body"]),
                Paragraph(str(item.confidence), styles["body"]),
            ]
        )
    if len(anomaly_rows) == 1:
        anomaly_rows.append(
            [Paragraph("冻结证据中未识别到结构化异常记录", styles["body"]), "—", "—", "—"]
        )
    anomaly_table = LongTable(
        anomaly_rows,
        colWidths=[39 * mm, 20 * mm, 91 * mm, 20 * mm],
        repeatRows=1,
    )
    anomaly_table.setStyle(_pdf_table_style())
    story.append(anomaly_table)

    story.append(Paragraph("行动建议", styles["heading"]))
    action_rows = [
        [
            Paragraph("优先级", styles["tableHeader"]),
            Paragraph("行动", styles["tableHeader"]),
            Paragraph("负责人", styles["tableHeader"]),
            Paragraph("截止时间", styles["tableHeader"]),
        ]
    ]
    for item in report.management_summary.actions:
        action_rows.append(
            [
                Paragraph(item.priority.upper(), styles["small"]),
                Paragraph(_pdf_text(item.text), styles["body"]),
                Paragraph(_pdf_text(item.owner or "待指派"), styles["body"]),
                Paragraph(item.due_at.strftime("%Y-%m-%d") if item.due_at else "待确认", styles["body"]),
            ]
        )
    for item in report.recommendations:
        action_rows.append(
            [
                Paragraph(item.priority.upper(), styles["small"]),
                Paragraph(_pdf_text(f"{item.title}：{item.rationale}"), styles["body"]),
                Paragraph("待指派", styles["body"]),
                Paragraph("待确认", styles["body"]),
            ]
        )
    action_table = LongTable(action_rows, colWidths=[22 * mm, 88 * mm, 32 * mm, 28 * mm], repeatRows=1)
    action_table.setStyle(_pdf_table_style())
    story.extend([action_table, Spacer(1, 5 * mm)])
    story.append(
        Paragraph(
            _pdf_text("数据来源：" + "、".join(report.appendix.sources)) + "。本文件由 ForgeSight 服务端基于冻结版本生成。",
            styles["small"],
        )
    )
    document.build(story, onFirstPage=_page_decoration, onLaterPages=_page_decoration)
    return output.getvalue()


def _pdf_table_style() -> TableStyle:
    return TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), BRAND_DARK),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, LINE),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, SURFACE]),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]
    )


def _prepare_sheet(sheet, title: str, report: StructuredWeeklyReport, headers: list[str]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(1, 1, _excel_safe(title))
    title_cell.font = Font(name="Microsoft YaHei", size=18, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="084C49")
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.cell(3, 1, "报告周期")
    sheet.cell(3, 2, report.identity.period_start)
    sheet.cell(3, 3, "至")
    sheet.cell(3, 4, report.identity.period_end)
    sheet.cell(3, 2).number_format = "yyyy-mm-dd"
    sheet.cell(3, 4).number_format = "yyyy-mm-dd"
    sheet.cell(4, 1, "报告版本")
    sheet.cell(4, 2, report.identity.version_id)
    sheet.cell(4, 3, "导出说明")
    sheet.cell(4, 4, "数据来自同一已冻结周报版本")
    sheet.freeze_panes = "A6"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:6"
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(6, column, value)
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B6E69")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.auto_filter.ref = f"A6:{get_column_letter(len(headers))}6"


def _finish_sheet(sheet, widths: list[int]) -> None:
    thin = Side(style="thin", color="D9E0E8")
    for row in sheet.iter_rows(min_row=6):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row > 6:
                cell.font = Font(name="Microsoft YaHei", size=10, color="172033")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _excel_safe(value):
    """Prevent spreadsheet applications from executing untrusted report text."""

    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _style_excel_chart(chart, spec: ChartSpec) -> None:
    chart.title = spec.title
    chart.height = 7.2
    chart.width = 12.8
    chart.style = 13
    chart.legend.position = "b"
    chart.graphical_properties = None
    colors_by_role = {
        "production": "0B6E69",
        "quality": "D88A16",
        "equipment": "7357C8",
        "inventory": "2B8A9A",
        "neutral": "667085",
    }
    colors_by_accent = {
        "ocean": "25677A",
        "teal": "0B6E69",
        "violet": "7357C8",
        "graphite": "58636D",
    }
    accent = colors_by_accent.get(spec.accent, colors_by_role.get(spec.color_role, "667085"))
    for series in chart.series:
        try:
            series.graphicalProperties.solidFill = accent
            series.graphicalProperties.line.solidFill = accent
            if spec.type == "scatter":
                series.graphicalProperties.line.noFill = True
                series.marker.symbol = "circle"
                series.marker.size = 7
        except AttributeError:
            continue


def _native_excel_chart(sheet, spec: ChartSpec, header_row: int, first_row: int, last_row: int):
    categories = Reference(sheet, min_col=1, min_row=first_row, max_row=last_row)
    primary = Reference(sheet, min_col=3, min_row=header_row, max_row=last_row)
    secondary = Reference(sheet, min_col=4, min_row=header_row, max_row=last_row)
    has_previous = any(point.previous_value is not None for point in spec.points)
    previous = Reference(sheet, min_col=7, min_row=header_row, max_row=last_row)
    if spec.type == "donut":
        chart = DoughnutChart()
        chart.add_data(primary, titles_from_data=True)
        chart.set_categories(categories)
        chart.holeSize = 55
    elif spec.type == "scatter":
        chart = ScatterChart()
        chart.x_axis.title = f"{spec.value_label}（{spec.unit}）"
        chart.y_axis.title = f"{spec.secondary_value_label or '次指标'}（{spec.secondary_unit or spec.unit}）"
        x_values = Reference(sheet, min_col=3, min_row=first_row, max_row=last_row)
        y_values = Reference(sheet, min_col=4, min_row=first_row, max_row=last_row)
        chart.series.append(Series(y_values, x_values, title=spec.secondary_value_label or "关系"))
    elif spec.type == "area":
        chart = AreaChart()
        chart.add_data(primary, titles_from_data=True)
        if has_previous:
            chart.add_data(previous, titles_from_data=True)
        chart.set_categories(categories)
    elif spec.type == "line":
        chart = LineChart()
        chart.add_data(primary, titles_from_data=True)
        if has_previous:
            chart.add_data(previous, titles_from_data=True)
        chart.set_categories(categories)
    elif spec.type == "column":
        chart = BarChart()
        chart.type = "col"
        chart.add_data(primary, titles_from_data=True)
        if has_previous:
            chart.add_data(previous, titles_from_data=True)
        chart.set_categories(categories)
    elif spec.type in {"combo", "pareto"}:
        chart = BarChart()
        chart.type = "col"
        chart.add_data(primary, titles_from_data=True)
        chart.set_categories(categories)
        line = LineChart()
        line.add_data(secondary, titles_from_data=True)
        line.y_axis.axId = 200
        line.y_axis.crosses = "max"
        chart += line
    else:
        chart = BarChart()
        chart.type = "bar"
        chart.add_data(primary, titles_from_data=True)
        chart.set_categories(categories)
    _style_excel_chart(chart, spec)
    return chart


def _create_chart_sheet(workbook: Workbook, report: StructuredWeeklyReport) -> None:
    sheet = workbook.create_sheet("图表数据")
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F1")
    sheet["A1"] = _excel_safe(report.identity.title + " · 图表数据")
    sheet["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="084C49")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet["A3"] = "报告周期"
    sheet["B3"] = report.identity.period_start
    sheet["C3"] = "至"
    sheet["D3"] = report.identity.period_end
    sheet["B3"].number_format = "yyyy-mm-dd"
    sheet["D3"].number_format = "yyyy-mm-dd"
    sheet["A4"] = "报告版本"
    sheet["B4"] = report.identity.version_id
    sheet["C4"] = "导出说明"
    sheet["D4"] = "所有图表与站内周报使用同一冻结版本"
    sheet.freeze_panes = "A8"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 18
    sheet.column_dimensions["E"].width = 34
    sheet.column_dimensions["F"].width = 14
    sheet.column_dimensions["G"].width = 18
    sheet.column_dimensions["H"].width = 2

    specs = _export_chart_specs(report)
    if not specs:
        sheet.merge_cells("A7:F9")
        sheet["A7"] = "当前冻结版本没有可验证的图表数据。"
        sheet["A7"].alignment = Alignment(horizontal="center", vertical="center")
        sheet["A7"].font = Font(name="Microsoft YaHei", color="667085")
        return

    row = 6
    thin = Side(style="thin", color="D9E0E8")
    for spec in specs:
        title_row = row
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        title_cell = sheet.cell(row, 1, _excel_safe(spec.title))
        title_cell.font = Font(name="Microsoft YaHei", size=13, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor="0B6E69")
        title_cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[row].height = 25
        row += 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        sheet.cell(row, 1, _excel_safe(spec.description)).font = Font(
            name="Microsoft YaHei", size=9, color="667085"
        )
        row += 1
        header_row = row
        has_previous = any(point.previous_value is not None for point in spec.points)
        secondary_header = (
            "累计贡献率（%）"
            if spec.type == "pareto"
            else f"{spec.secondary_value_label or '次指标'}（{spec.secondary_unit or spec.unit}）"
        )
        headers = [
            spec.category_label,
            "分组",
            f"{spec.value_label}（{spec.unit}）",
            secondary_header,
            "证据引用",
            "图表类型",
        ]
        if has_previous:
            headers.append(f"上一周期（{spec.unit}）")
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row, column, _excel_safe(value))
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="172033")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row += 1
        first_data_row = row
        for point in spec.points:
            values = [
                point.label,
                point.group,
                point.value,
                point.secondary_value,
                "、".join(spec.evidence_artifact_ids),
                spec.type,
            ]
            if has_previous:
                values.append(point.previous_value)
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row, column, _excel_safe(value))
                cell.font = Font(name="Microsoft YaHei", size=9, color="172033")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
            row += 1
        last_data_row = row - 1
        if spec.type == "heatmap":
            sheet.conditional_formatting.add(
                f"C{first_data_row}:C{last_data_row}",
                ColorScaleRule(
                    start_type="min", start_color="DDF4EF",
                    mid_type="percentile", mid_value=50, mid_color="7FCFC4",
                    end_type="max", end_color="0B6E69",
                ),
            )
            groups = list(dict.fromkeys(point.group for point in spec.points if point.group))
            labels = list(dict.fromkeys(point.label for point in spec.points))
            values = {(point.label, point.group): point.value for point in spec.points}
            pivot_row = title_row
            pivot_column = 8
            sheet.cell(pivot_row, pivot_column, _excel_safe(spec.title)).font = Font(
                name="Microsoft YaHei", size=11, bold=True, color="172033"
            )
            for index, group in enumerate(groups, start=1):
                cell = sheet.cell(pivot_row + 1, pivot_column + index, _excel_safe(group))
                cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="084C49")
                cell.alignment = Alignment(horizontal="center")
            for label_index, label in enumerate(labels, start=1):
                label_cell = sheet.cell(pivot_row + 1 + label_index, pivot_column, _excel_safe(label))
                label_cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color="172033")
                label_cell.fill = PatternFill("solid", fgColor="F5F7FA")
                for group_index, group in enumerate(groups, start=1):
                    value_cell = sheet.cell(
                        pivot_row + 1 + label_index,
                        pivot_column + group_index,
                        values.get((label, group)),
                    )
                    value_cell.number_format = "0.00"
                    value_cell.alignment = Alignment(horizontal="center")
            if labels and groups:
                heatmap_range = (
                    f"{get_column_letter(pivot_column + 1)}{pivot_row + 2}:"
                    f"{get_column_letter(pivot_column + len(groups))}{pivot_row + 1 + len(labels)}"
                )
                sheet.conditional_formatting.add(
                    heatmap_range,
                    ColorScaleRule(
                        start_type="min", start_color="DDF4EF",
                        mid_type="percentile", mid_value=50, mid_color="7FCFC4",
                        end_type="max", end_color="0B6E69",
                    ),
                )
        elif spec.type == "calendar":
            parsed = []
            for point in spec.points:
                day = _parse_calendar_date(point.label)
                if day is not None:
                    parsed.append((day, point.value or 0))
            if len(parsed) >= 2:
                from datetime import timedelta
                parsed.sort(key=lambda item: item[0])
                grid_start = parsed[0][0] - timedelta(days=parsed[0][0].weekday())
                last_day = parsed[-1][0]
                weeks = ((last_day - grid_start).days + 7) // 7
                by_day = dict(parsed)
                pivot_row = title_row
                pivot_column = 8
                sheet.cell(pivot_row, pivot_column, _excel_safe(spec.title)).font = Font(
                    name="Microsoft YaHei", size=11, bold=True, color="172033"
                )
                for index, label in enumerate(["一", "二", "三", "四", "五", "六", "日"]):
                    head = sheet.cell(pivot_row + 1, pivot_column + index, label)
                    head.font = Font(name="Microsoft YaHei", size=9, bold=True, color="FFFFFF")
                    head.fill = PatternFill("solid", fgColor="084C49")
                    head.alignment = Alignment(horizontal="center")
                for week in range(weeks):
                    for weekday in range(7):
                        current = grid_start + timedelta(days=week * 7 + weekday)
                        cell = sheet.cell(pivot_row + 2 + week, pivot_column + weekday)
                        if current in by_day:
                            cell.value = by_day[current]
                            cell.number_format = "0"
                        cell.alignment = Alignment(horizontal="center")
                grid_range = (
                    f"{get_column_letter(pivot_column)}{pivot_row + 2}:"
                    f"{get_column_letter(pivot_column + 6)}{pivot_row + 1 + weeks}"
                )
                sheet.conditional_formatting.add(
                    grid_range,
                    ColorScaleRule(
                        start_type="min", start_color="DDF4EF",
                        mid_type="percentile", mid_value=50, mid_color="7FCFC4",
                        end_type="max", end_color="0B6E69",
                    ),
                )
        elif spec.type != "table":
            native_chart = _native_excel_chart(
                sheet, spec, header_row, first_data_row, last_data_row
            )
            sheet.add_chart(native_chart, f"H{title_row}")
        row = max(row + 2, title_row + 17)
    sheet.print_area = f"A1:S{row}"


def render_excel_bytes(content: dict) -> bytes:
    """Render a typed, styled workbook from the same frozen report document."""

    report = StructuredWeeklyReport.model_validate(content)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "指标概览"
    _prepare_sheet(
        overview,
        report.identity.title + " · 指标概览",
        report,
        ["指标", "本期值", "单位", "上期值", "环比", "可信度"],
    )
    for row_index, metric in enumerate(report.metrics, start=7):
        values = [
            metric.name,
            metric.value,
            metric.unit,
            metric.previous_value,
            metric.change_rate,
            metric.confidence,
        ]
        for column, value in enumerate(values, start=1):
            overview.cell(row_index, column, _excel_safe(value))
        overview.cell(row_index, 5).number_format = "0.00%"
    _finish_sheet(overview, [24, 16, 12, 16, 14, 12])

    _create_chart_sheet(workbook, report)

    anomalies = workbook.create_sheet("异常明细")
    _prepare_sheet(
        anomalies,
        report.identity.title + " · 异常明细",
        report,
        ["异常", "等级", "指标", "说明", "可信度", "证据引用"],
    )
    for row_index, item in enumerate(report.anomalies, start=7):
        values = [
            item.title,
            item.severity,
            item.metric_key,
            item.description,
            item.confidence,
            "、".join(item.evidence_artifact_ids),
        ]
        for column, value in enumerate(values, start=1):
            anomalies.cell(row_index, column, _excel_safe(value))
    _finish_sheet(anomalies, [28, 12, 18, 52, 12, 30])

    actions = workbook.create_sheet("行动建议")
    _prepare_sheet(
        actions,
        report.identity.title + " · 行动建议",
        report,
        ["优先级", "行动", "负责人", "截止时间", "依据"],
    )
    row_index = 7
    for item in report.management_summary.actions:
        values = [
            item.priority,
            item.text,
            item.owner,
            item.due_at,
            "、".join(item.evidence_artifact_ids),
        ]
        for column, value in enumerate(values, start=1):
            actions.cell(row_index, column, _excel_safe(value))
        actions.cell(row_index, 4).number_format = "yyyy-mm-dd hh:mm"
        row_index += 1
    for item in report.recommendations:
        values = [item.priority, f"{item.title}：{item.rationale}", None, None, "、".join(item.evidence_artifact_ids)]
        for column, value in enumerate(values, start=1):
            actions.cell(row_index, column, _excel_safe(value))
        row_index += 1
    _finish_sheet(actions, [12, 58, 20, 22, 30])

    sources = workbook.create_sheet("数据来源与口径")
    _prepare_sheet(
        sources,
        report.identity.title + " · 数据来源与口径",
        report,
        ["类型", "名称", "说明", "Schema 版本"],
    )
    row_index = 7
    for source in report.appendix.sources:
        sources.append(["数据来源", _excel_safe(source), "已冻结版本引用", report.schema_version])
        row_index += 1
    for definition in report.appendix.definitions:
        sources.append(
            [
                "指标口径",
                _excel_safe(definition.name),
                _excel_safe(definition.definition),
                report.schema_version,
            ]
        )
    _finish_sheet(sources, [16, 30, 70, 24])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
